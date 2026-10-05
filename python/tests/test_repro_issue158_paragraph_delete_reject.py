"""
Regression tests for GitHub Issue #158:
"Deleting a paragraph with its separator merges the next paragraph into it; Reject All cannot restore the break"

Cross-platform parity: matching tests live in
    node/packages/core/src/engine.bugs.test.ts (Issue #158 test)
"""

import io

from docx import Document

from adeu.models import ModifyText, RejectChange
from adeu.redline.engine import RedlineEngine


def test_issue158_paragraph_deletion_reject_all_restores_break():
    doc = Document()
    for text in ("First paragraph.", "Second paragraph.", "Third paragraph."):
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    initial_bytes = buf.getvalue()

    engine = RedlineEngine(io.BytesIO(initial_bytes), author="Tester")
    stats = engine.process_batch(
        [ModifyText(type="modify", target_text="Second paragraph.\n\n", new_text="", comment=None)]
    )
    assert stats["status"] == "ok"
    assert stats["edits"][0]["status"] == "applied"
    out = engine.save_to_stream().getvalue()

    # Accept all must cleanly leave only First and Third
    acc_engine = RedlineEngine(io.BytesIO(out), author="Tester")
    acc_engine.accept_all_revisions()
    acc_doc = Document(io.BytesIO(acc_engine.save_to_stream().getvalue()))
    assert [p.text for p in acc_doc.paragraphs] == ["First paragraph.", "Third paragraph."]

    # Reject all must cleanly restore all three separate paragraphs
    rej_engine = RedlineEngine(io.BytesIO(out), author="Tester")
    rej_engine.reject_all_revisions()
    rej_doc = Document(io.BytesIO(rej_engine.save_to_stream().getvalue()))
    assert [p.text for p in rej_doc.paragraphs] == [
        "First paragraph.",
        "Second paragraph.",
        "Third paragraph.",
    ]


def test_issue158_paragraph_deletion_individual_reject_restores_break():
    doc = Document()
    for text in ("First paragraph.", "Second paragraph.", "Third paragraph."):
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    initial_bytes = buf.getvalue()

    engine = RedlineEngine(io.BytesIO(initial_bytes), author="Tester")
    engine.process_batch([ModifyText(type="modify", target_text="Second paragraph.\n\n", new_text="", comment=None)])
    out = engine.save_to_stream().getvalue()

    # Reject via apply_review_actions on all generated deletion IDs
    rev_engine = RedlineEngine(io.BytesIO(out), author="Tester")
    del_ids = [
        d.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id")
        for d in rev_engine._story_findall("w:del")
    ]
    for did in set(filter(None, del_ids)):
        rev_engine.apply_review_actions([RejectChange(type="reject", target_id=f"Chg:{did}", part=None, comment=None)])

    rev_doc = Document(io.BytesIO(rev_engine.save_to_stream().getvalue()))
    assert [p.text for p in rev_doc.paragraphs] == [
        "First paragraph.",
        "Second paragraph.",
        "Third paragraph.",
    ]
