import io

import pytest
from docx import Document

from adeu.models import ModifyText
from adeu.redline.engine import BatchValidationError, RedlineEngine


def test_cross_table_cell_edit_validation_error():
    """
    Test Case: When an edit tries to replace text that spans across table cells
    (which ingest formats with ` | `), and the new_text does not have the same
    number of separators, the engine should raise a BatchValidationError.
    """
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "CellA"
    table.cell(0, 1).text = "CellB"

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    # Ingest output: "CellA | CellB"
    edit = ModifyText(target_text="CellA | CellB", new_text="CellC")
    engine = RedlineEngine(stream)

    with pytest.raises(BatchValidationError) as exc_info:
        engine.apply_edits([edit])

    assert "Target text spans 2 table cells, but replacement provides 1" in str(exc_info.value)
