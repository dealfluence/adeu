"""
Regression tests for GitHub Issue #157:
"A tab in a run that an edit splits is replaced by a space, untracked"

Cross-platform parity: matching tests live in
    node/packages/core/src/engine.bugs.test.ts (Issue #157 test)
"""

import io
import re
import zipfile

from docx import Document

from adeu.models import ModifyText
from adeu.redline.engine import RedlineEngine


def _body_xml(docx_bytes: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    m = re.search(r"<w:body>(.*?)<w:sectPr", xml, re.DOTALL)
    return m.group(1) if m else xml


def test_issue157_tab_preserved_on_run_split():
    doc = Document()
    doc.add_paragraph("1.1\tThe Supplier shall deliver the Goods.")
    buf = io.BytesIO()
    doc.save(buf)
    initial_bytes = buf.getvalue()

    tabs_before = _body_xml(initial_bytes).count("<w:tab/>")
    assert tabs_before == 1

    engine = RedlineEngine(io.BytesIO(initial_bytes), author="Tester")
    engine.apply_edits([ModifyText(type="modify", target_text="deliver", new_text="supply", comment=None)])
    out = engine.save_to_stream().getvalue()

    tabs_after = _body_xml(out).count("<w:tab/>")
    assert tabs_after == 1, f"Expected 1 <w:tab/>, found {tabs_after}"

    # Rejecting revisions must restore original text with tab intact
    rejected_engine = RedlineEngine(io.BytesIO(out), author="Tester")
    rejected_engine.reject_all_revisions()
    rejected_doc = Document(io.BytesIO(rejected_engine.save_to_stream().getvalue()))
    assert rejected_doc.paragraphs[0].text == "1.1\tThe Supplier shall deliver the Goods."

    # Accepting revisions must keep tab intact
    accepted_engine = RedlineEngine(io.BytesIO(out), author="Tester")
    accepted_engine.accept_all_revisions()
    accepted_doc = Document(io.BytesIO(accepted_engine.save_to_stream().getvalue()))
    assert accepted_doc.paragraphs[0].text == "1.1\tThe Supplier shall supply the Goods."


def test_issue157_tab_preserved_when_adjacent_words_edited():
    doc = Document()
    doc.add_paragraph("1.1\tThe Supplier shall deliver the Goods.")
    buf = io.BytesIO()
    doc.save(buf)
    initial_bytes = buf.getvalue()

    # Edit "The" immediately after the tab
    engine = RedlineEngine(io.BytesIO(initial_bytes), author="Tester")
    engine.apply_edits([ModifyText(type="modify", target_text="The", new_text="A", comment=None)])
    out = engine.save_to_stream().getvalue()

    assert _body_xml(out).count("<w:tab/>") == 1

    rejected_engine = RedlineEngine(io.BytesIO(out), author="Tester")
    rejected_engine.reject_all_revisions()
    rejected_doc = Document(io.BytesIO(rejected_engine.save_to_stream().getvalue()))
    assert rejected_doc.paragraphs[0].text == "1.1\tThe Supplier shall deliver the Goods."

    accepted_engine = RedlineEngine(io.BytesIO(out), author="Tester")
    accepted_engine.accept_all_revisions()
    accepted_doc = Document(io.BytesIO(accepted_engine.save_to_stream().getvalue()))
    assert accepted_doc.paragraphs[0].text == "1.1\tA Supplier shall deliver the Goods."
