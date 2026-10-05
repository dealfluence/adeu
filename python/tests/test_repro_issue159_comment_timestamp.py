"""
Regression tests for GitHub Issue #159:
"Revision and comment timestamps dated from different clocks in Python and Node engines."

Cross-platform parity: matching tests live in
    node/packages/core/src/engine.bugs.test.ts (Issue #159 tests)
"""

import datetime
import io
import re
from typing import Any

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from adeu.models import ModifyText, ReplyComment
from adeu.redline.comments import CommentsManager
from adeu.redline.engine import RedlineEngine


def _get_comments_xml(doc: Any) -> str:
    for rel in doc.part.rels.values():
        if rel.reltype == RT.COMMENTS:
            return rel.target_part.blob.decode("utf-8")
    return ""


def test_issue159_revisions_and_comments_share_engine_timestamp():
    doc = Document()
    doc.add_paragraph("The quick brown fox.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    fixed_time = "2026-04-15T09:30:00Z"
    engine = RedlineEngine(buf, author="Tester", timestamp=fixed_time)
    assert engine.timestamp == fixed_time

    edit = ModifyText(type="modify", target_text="quick", new_text="slow", comment="Foxes are not always quick.")
    stats = engine.process_batch([edit])
    assert stats["status"] == "ok"

    res_stream = engine.save_to_stream()
    res_doc = Document(res_stream)

    # Document body XML: revision elements should have fixed_time
    doc_xml = res_doc.element.xml
    del_dates = re.findall(r'<w:del[^>]*w:date="([^"]+)"', doc_xml)
    ins_dates = re.findall(r'<w:ins[^>]*w:date="([^"]+)"', doc_xml)
    assert del_dates or ins_dates
    for d in del_dates + ins_dates:
        assert d == fixed_time

    # Comments XML: w:comment w:date should also have fixed_time
    comments_xml = _get_comments_xml(res_doc)
    assert comments_xml
    comment_dates = re.findall(r'<w:comment[^>]*w:date="([^"]+)"', comments_xml)
    assert comment_dates
    for d in comment_dates:
        assert d == fixed_time


def test_issue159_engine_init_timestamp_types():
    doc = Document()
    doc.add_paragraph("Sample text.")
    buf = io.BytesIO()
    doc.save(buf)

    # 1. str timestamp
    eng_str = RedlineEngine(io.BytesIO(buf.getvalue()), timestamp="2026-01-01T00:00:00Z")
    assert eng_str.timestamp == "2026-01-01T00:00:00Z"

    # 2. Naive datetime -> converted to UTC format
    dt_naive = datetime.datetime(2026, 6, 1, 12, 0, 0)
    eng_naive = RedlineEngine(io.BytesIO(buf.getvalue()), timestamp=dt_naive)
    assert eng_naive.timestamp == "2026-06-01T12:00:00Z"

    # 3. Aware datetime with offset -> converted to UTC
    tz_plus_2 = datetime.timezone(datetime.timedelta(hours=2))
    dt_aware = datetime.datetime(2026, 6, 1, 14, 0, 0, tzinfo=tz_plus_2)
    eng_aware = RedlineEngine(io.BytesIO(buf.getvalue()), timestamp=dt_aware)
    assert eng_aware.timestamp == "2026-06-01T12:00:00Z"

    # 4. None -> defaults to current UTC formatted string
    eng_none = RedlineEngine(io.BytesIO(buf.getvalue()), timestamp=None)
    assert re.match(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z$", eng_none.timestamp)


def test_issue159_reassign_engine_timestamp_before_process_batch():
    doc = Document()
    doc.add_paragraph("The quick brown fox.")
    buf = io.BytesIO()
    doc.save(buf)

    engine = RedlineEngine(io.BytesIO(buf.getvalue()), author="Tester")
    reassigned_time = "2026-12-25T18:00:00Z"
    engine.timestamp = reassigned_time

    edit = ModifyText(type="modify", target_text="quick", new_text="slow", comment="Foxes are slow.")
    stats = engine.process_batch([edit])
    assert stats["status"] == "ok"

    res_doc = Document(engine.save_to_stream())
    comments_xml = _get_comments_xml(res_doc)
    comment_dates = re.findall(r'<w:comment[^>]*w:date="([^"]+)"', comments_xml)
    assert comment_dates == [reassigned_time]

    doc_xml = res_doc.element.xml
    del_dates = re.findall(r'<w:del[^>]*w:date="([^"]+)"', doc_xml)
    ins_dates = re.findall(r'<w:ins[^>]*w:date="([^"]+)"', doc_xml)
    for d in del_dates + ins_dates:
        assert d == reassigned_time


def test_issue159_reply_comment_uses_engine_timestamp():
    doc = Document()
    doc.add_paragraph("The quick brown fox.")
    buf = io.BytesIO()
    doc.save(buf)

    engine = RedlineEngine(io.BytesIO(buf.getvalue()), author="Tester", timestamp="2026-01-01T10:00:00Z")
    engine.process_batch([ModifyText(type="modify", target_text="quick", new_text="fast", comment="Original note")])

    # Reply with a new timestamp
    reply_time = "2026-01-02T15:00:00Z"
    engine.timestamp = reply_time
    applied, skipped, already_resolved = engine.apply_review_actions(
        [ReplyComment(type="reply", target_id="Com:1", text="Reply note")]
    )
    assert applied == 1

    res_doc = Document(engine.save_to_stream())
    comments_xml = _get_comments_xml(res_doc)
    comments = re.findall(r'<w:comment[^>]*w:id="(\d+)"[^>]*w:date="([^"]+)"', comments_xml)
    comment_map = {cid: dt for cid, dt in comments}
    assert comment_map["1"] == "2026-01-01T10:00:00Z"
    assert comment_map["2"] == reply_time


def test_issue159_comments_manager_direct_timestamp():
    doc = Document()
    cm = CommentsManager(doc)
    cid1 = cm.add_comment("Tester", "Explicit timestamp", timestamp="2026-07-07T07:07:07Z")
    cid2 = cm.add_comment("Tester", "Default timestamp", timestamp=None)

    data = cm.extract_comments_data()
    assert data[cid1]["date"] == "2026-07-07T07:07:07Z"
    assert re.match(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z$", data[cid2]["date"])
