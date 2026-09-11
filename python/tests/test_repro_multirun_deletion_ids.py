"""Synthetic regression for duplicate w:del IDs across runs in one edit."""

import io

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from adeu.ingest import extract_text_from_stream
from adeu.models import AcceptChange, ModifyText, RejectChange
from adeu.redline.engine import RedlineEngine


@pytest.mark.parametrize("with_boundary", [False, True])
@pytest.mark.parametrize("replacement", ["violet", ""])
def test_multirun_deletion_ids_and_resolution(with_boundary: bool, replacement: str) -> None:
    doc = Document()
    paragraph = doc.add_paragraph("Before ")
    for index, text in enumerate(["red ", "green ", "blue"]):
        if with_boundary and index != 0:
            marker = OxmlElement("w:proofErr")
            marker.set(qn("w:type"), "spellStart" if index == 1 else "spellEnd")
            paragraph._p.append(marker)
        paragraph.add_run(text).font.size = Pt(12 + index)
    doc.add_paragraph("Untouched paragraph.")
    original = io.BytesIO()
    doc.save(original)
    original.seek(0)
    original_text = extract_text_from_stream(original, clean_view=True)
    original.seek(0)
    engine = RedlineEngine(original, author="Reviewer")
    report = engine.process_batch([ModifyText(target_text="red green blue", new_text=replacement)])
    assert report["edits_applied"] == 1
    edited = engine.save_to_stream().getvalue()
    reloaded = Document(io.BytesIO(edited))
    ids = [node.get(qn("w:id")) for node in reloaded.element.xpath("//w:del")]
    assert len(ids) == (3 if with_boundary else 1)
    assert len(set(ids)) == len(ids)
    assert [node.get(qn("w:val")) for node in reloaded.element.xpath("//w:del//w:sz")] == ["24", "26", "28"]
    assert len(reloaded.element.xpath("//w:proofErr")) == (2 if with_boundary else 0)
    assert all(node.getparent().tag == qn("w:p") for node in reloaded.element.xpath("//w:proofErr"))
    if replacement:
        assert reloaded.element.xpath("//w:ins//w:sz")[0].get(qn("w:val")) == "28"
    for action in [AcceptChange, RejectChange]:
        resolved = RedlineEngine(io.BytesIO(edited), author="Reviewer")
        resolved.process_batch([action(target_id=target_id) for target_id in ids])
        text = extract_text_from_stream(resolved.save_to_stream(), clean_view=True)
        assert text == (
            original_text if action == RejectChange else original_text.replace("red green blue", replacement)
        )
        result = Document(resolved.save_to_stream())
        assert not result.element.xpath("//w:del | //w:ins")


def test_preserves_foreign_revisions_after_reload() -> None:
    doc = Document()
    p = doc.add_paragraph("Before ")
    p.add_run("red ")
    p.add_run("green ")
    p.add_run("blue")
    doc.add_paragraph("Untouched paragraph.")
    original = io.BytesIO()
    doc.save(original)
    original.seek(0)
    engine = RedlineEngine(original, author="Other reviewer")
    engine.process_batch([ModifyText(target_text="Untouched", new_text="Protected")])
    reloaded = RedlineEngine(engine.save_to_stream(), author="Reviewer")
    foreign_paragraph = reloaded.doc.paragraphs[1]._p.xml
    reloaded.process_batch([ModifyText(target_text="red green blue", new_text="violet")])
    assert reloaded.doc.paragraphs[1]._p.xml == foreign_paragraph
    ids = [node.get(qn("w:id")) for node in reloaded.doc.element.xpath("//w:del")]
    assert len(set(ids)) == len(ids)
