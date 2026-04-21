"""
Regression tests for _build_mock_docx_stream.

Validates the Flat OPC -> ZIP-based DOCX reconstruction logic used by the live Word
read path. These tests are pure-Python and run on every platform (no COM required).
"""

import base64
import io
import zipfile

from docx import Document

from adeu.mcp_components.tools.live_word import _build_mock_docx_stream


def _wrap_pkg(parts_xml: str) -> str:
    """Wrap a set of pkg:part elements into a full Flat OPC envelope."""
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        '<?mso-application progid="Word.Document"?>\r\n'
        '<pkg:package xmlns:pkg="http://schemas.microsoft.com/office/2006/xmlPackage">\r\n'
        f"{parts_xml}\r\n"
        "</pkg:package>"
    )


_MINIMAL_DOCUMENT_XML = (
    '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    "<w:body><w:p><w:r><w:t>Hello world</w:t></w:r></w:p><w:sectPr/></w:body>"
    "</w:document>"
)

_MINIMAL_ROOT_RELS = (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
    'Target="word/document.xml"/>'
    "</Relationships>"
)


def _minimal_document_part() -> str:
    return (
        '<pkg:part pkg:name="/word/document.xml" '
        'pkg:contentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml">'
        f"<pkg:xmlData>{_MINIMAL_DOCUMENT_XML}</pkg:xmlData>"
        "</pkg:part>"
    )


def _minimal_rels_part() -> str:
    return (
        '<pkg:part pkg:name="/_rels/.rels" '
        'pkg:contentType="application/vnd.openxmlformats-package.relationships+xml">'
        f"<pkg:xmlData>{_MINIMAL_ROOT_RELS}</pkg:xmlData>"
        "</pkg:part>"
    )


def _open_zip(stream: io.BytesIO) -> zipfile.ZipFile:
    stream.seek(0)
    return zipfile.ZipFile(stream, "r")


def test_round_trips_through_python_docx():
    """The rebuilt archive must open cleanly via python-docx and expose the text."""
    flat = _wrap_pkg(_minimal_rels_part() + _minimal_document_part())
    stream = _build_mock_docx_stream(flat)

    stream.seek(0)
    doc = Document(stream)
    assert any("Hello world" in p.text for p in doc.paragraphs)


def test_synthesizes_content_types_xml_with_overrides():
    """[Content_Types].xml must carry an <Override> for every non-rels part written."""
    flat = _wrap_pkg(_minimal_rels_part() + _minimal_document_part())
    stream = _build_mock_docx_stream(flat)

    with _open_zip(stream) as zf:
        assert "[Content_Types].xml" in zf.namelist()
        ct = zf.read("[Content_Types].xml").decode("utf-8")

    # Single rels Default; no duplicate Override for the rels part
    assert 'Extension="rels"' in ct
    assert 'PartName="/_rels/.rels"' not in ct
    assert 'PartName="/word/document.xml"' in ct


def test_parses_self_closing_pkg_part_without_crashing():
    """
    Self-closing <pkg:part ... /> elements are legitimate (Word emits them for
    empty placeholders). The builder must accept them without raising; since
    there is no body, no archive entry is written and no Content Types Override
    is emitted — this mirrors the current documented policy.
    """
    empty_part = '<pkg:part pkg:name="/word/custom.xml" pkg:contentType="application/xml" />'
    flat = _wrap_pkg(_minimal_rels_part() + _minimal_document_part() + empty_part)
    stream = _build_mock_docx_stream(flat)

    with _open_zip(stream) as zf:
        names = zf.namelist()
        ct = zf.read("[Content_Types].xml").decode("utf-8")

    assert "word/custom.xml" not in names
    assert 'PartName="/word/custom.xml"' not in ct
    # Main parts still made it through
    assert "word/document.xml" in names


def test_parses_binary_parts_via_base64():
    """Parts with <pkg:binaryData> must be base64-decoded and written as raw bytes."""
    png_bytes = b"\x89PNG\r\n\x1a\n" + b"\x00" * 10
    b64 = base64.b64encode(png_bytes).decode("ascii")
    binary_part = (
        '<pkg:part pkg:name="/word/media/image1.png" '
        'pkg:contentType="image/png" '
        'pkg:compression="store">'
        f"<pkg:binaryData>{b64}</pkg:binaryData>"
        "</pkg:part>"
    )
    flat = _wrap_pkg(_minimal_rels_part() + _minimal_document_part() + binary_part)
    stream = _build_mock_docx_stream(flat)

    with _open_zip(stream) as zf:
        assert "word/media/image1.png" in zf.namelist()
        assert zf.read("word/media/image1.png") == png_bytes
        ct = zf.read("[Content_Types].xml").decode("utf-8")

    assert 'PartName="/word/media/image1.png"' in ct
    assert 'ContentType="image/png"' in ct


def test_rels_default_covers_multiple_rels_files_without_override_duplication():
    """
    Every .rels file in the package must be written as an archive entry,
    but none should appear as an <Override> — the single Default covers them.
    """
    extra_rels_part = (
        '<pkg:part pkg:name="/word/_rels/document.xml.rels" '
        'pkg:contentType="application/vnd.openxmlformats-package.relationships+xml">'
        "<pkg:xmlData>"
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>'
        "</pkg:xmlData>"
        "</pkg:part>"
    )
    flat = _wrap_pkg(_minimal_rels_part() + _minimal_document_part() + extra_rels_part)
    stream = _build_mock_docx_stream(flat)

    with _open_zip(stream) as zf:
        names = zf.namelist()
        ct = zf.read("[Content_Types].xml").decode("utf-8")

    assert "_rels/.rels" in names
    assert "word/_rels/document.xml.rels" in names

    # Neither rels file gets an Override — the Default handles them
    assert 'PartName="/_rels/.rels"' not in ct
    assert 'PartName="/word/_rels/document.xml.rels"' not in ct
    # And only one Default for rels
    assert ct.count('Extension="rels"') == 1


def test_missing_name_or_content_type_is_skipped_without_crash():
    """Parts with malformed attribute strings must be skipped, not raise."""
    malformed_missing_ctype = '<pkg:part pkg:name="/foo.xml"><pkg:xmlData><x/></pkg:xmlData></pkg:part>'
    malformed_missing_name = '<pkg:part pkg:contentType="application/xml"><pkg:xmlData><x/></pkg:xmlData></pkg:part>'
    flat = _wrap_pkg(_minimal_rels_part() + _minimal_document_part() + malformed_missing_ctype + malformed_missing_name)
    # Must not raise
    stream = _build_mock_docx_stream(flat)

    with _open_zip(stream) as zf:
        names = zf.namelist()

    assert "foo.xml" not in names
    # But the well-formed parts survived the skip
    assert "word/document.xml" in names
    assert "_rels/.rels" in names
